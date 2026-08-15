import csv
import io
from datetime import datetime
from statistics import mean, stdev
from typing import List, Optional

from fastapi import HTTPException
from openpyxl import Workbook
from sqlalchemy import case, func, or_
from sqlalchemy.orm import Session

from app.Auth.utils import pwd_context
from app.admin import schemas
from app.models.tables import (
    Category,
    Create_Account_Table,
    Transaction,
    TransactionType,
)

ALLOWED_ROLES = ("user", "admin")


def _get_user(db: Session, user_id: int) -> Create_Account_Table:
    user = (
        db.query(Create_Account_Table)
        .filter(Create_Account_Table.profile_id == user_id)
        .first()
    )
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    return user


def _get_transaction(db: Session, transaction_id: int) -> Transaction:
    transaction = (
        db.query(Transaction).filter(Transaction.id == transaction_id).first()
    )
    if not transaction:
        raise HTTPException(status_code=404, detail="Transaction not found")
    return transaction


def _income_expr():
    return func.sum(
        case((Transaction.type == TransactionType.income, Transaction.amount), else_=0)
    ).label("total_income")


def _expense_expr():
    return func.sum(
        case((Transaction.type == TransactionType.expense, Transaction.amount), else_=0)
    ).label("total_expense")


def _build_user_stats(db: Session, users: List[Create_Account_Table]) -> List[schemas.UserStats]:
    ids = [u.profile_id for u in users]
    if not ids:
        return []

    rows = (
        db.query(
            Transaction.user_id,
            _income_expr(),
            _expense_expr(),
            func.count(Transaction.id).label("transaction_count"),
        )
        .filter(Transaction.user_id.in_(ids))
        .group_by(Transaction.user_id)
        .all()
    )
    stats_map = {row.user_id: row for row in rows}

    result = []
    for u in users:
        row = stats_map.get(u.profile_id)
        result.append(
            schemas.UserStats(
                profile_id=u.profile_id,
                name=u.name,
                email=u.email,
                mobile_number=u.mobile_number,
                role=u.role,
                is_active=u.is_active,
                created_at=u.created_at,
                total_income=float(row.total_income or 0) if row else 0.0,
                total_expense=float(row.total_expense or 0) if row else 0.0,
                transaction_count=int(row.transaction_count or 0) if row else 0,
            )
        )
    return result


def _to_admin_transaction(t: Transaction) -> schemas.AdminTransactionOut:
    return schemas.AdminTransactionOut(
        id=t.id,
        user_id=t.user_id,
        user_name=t.user.name,
        user_email=t.user.email,
        type=t.type,
        amount=t.amount,
        category=t.category,
        note=t.note,
        transaction_datetime=t.transaction_datetime,
        created_at=t.created_at,
        updated_at=t.updated_at,
    )


# ------------------------------ User management ------------------------------


def list_users(
    db: Session,
    search: Optional[str] = None,
    role: Optional[str] = None,
    is_active: Optional[bool] = None,
    skip: int = 0,
    limit: int = 50,
) -> schemas.UserListResponse:
    query = db.query(Create_Account_Table)
    if search:
        like = f"%{search}%"
        query = query.filter(
            or_(
                Create_Account_Table.name.ilike(like),
                Create_Account_Table.email.ilike(like),
                Create_Account_Table.mobile_number.ilike(like),
            )
        )
    if role:
        query = query.filter(Create_Account_Table.role == role)
    if is_active is not None:
        query = query.filter(Create_Account_Table.is_active.is_(is_active))

    total = query.count()
    users = (
        query.order_by(Create_Account_Table.profile_id)
        .offset(skip)
        .limit(limit)
        .all()
    )
    return schemas.UserListResponse(total=total, users=_build_user_stats(db, users))


def get_user_with_stats(db: Session, user_id: int) -> schemas.UserStats:
    user = _get_user(db, user_id)
    return _build_user_stats(db, [user])[0]


def set_user_status(db: Session, user_id: int, is_active: bool) -> schemas.UserStats:
    user = _get_user(db, user_id)
    user.is_active = is_active
    db.commit()
    db.refresh(user)
    return _build_user_stats(db, [user])[0]


def set_user_role(db: Session, user_id: int, role: str) -> schemas.UserStats:
    if role not in ALLOWED_ROLES:
        raise HTTPException(status_code=400, detail="Role must be 'user' or 'admin'")
    user = _get_user(db, user_id)
    user.role = role
    db.commit()
    db.refresh(user)
    return _build_user_stats(db, [user])[0]


def reset_password(db: Session, user_id: int, new_password: str) -> schemas.UserStats:
    user = _get_user(db, user_id)
    user.password = pwd_context.hash(new_password)
    user.token_version += 1
    db.commit()
    db.refresh(user)
    return _build_user_stats(db, [user])[0]


def force_logout(db: Session, user_id: int) -> schemas.UserStats:
    user = _get_user(db, user_id)
    user.token_version += 1
    db.commit()
    db.refresh(user)
    return _build_user_stats(db, [user])[0]


def delete_user(db: Session, user_id: int) -> schemas.UserStats:
    user = _get_user(db, user_id)
    stats = _build_user_stats(db, [user])[0]
    db.query(Transaction).filter(Transaction.user_id == user_id).delete()
    db.delete(user)
    db.commit()
    return stats


# --------------------------- Transaction oversight ---------------------------


def list_transactions(
    db: Session,
    user_id: Optional[int] = None,
    type: Optional[TransactionType] = None,
    category: Optional[str] = None,
    date_from: Optional[datetime] = None,
    date_to: Optional[datetime] = None,
    min_amount: Optional[int] = None,
    max_amount: Optional[int] = None,
    search: Optional[str] = None,
    skip: int = 0,
    limit: int = 50,
) -> schemas.AdminTransactionListResponse:
    query = db.query(Transaction).join(
        Create_Account_Table, Transaction.user_id == Create_Account_Table.profile_id
    )
    if user_id is not None:
        query = query.filter(Transaction.user_id == user_id)
    if type is not None:
        query = query.filter(Transaction.type == type)
    if category:
        query = query.filter(Transaction.category.ilike(f"%{category}%"))
    if date_from is not None:
        query = query.filter(Transaction.transaction_datetime >= date_from)
    if date_to is not None:
        query = query.filter(Transaction.transaction_datetime <= date_to)
    if min_amount is not None:
        query = query.filter(Transaction.amount >= min_amount)
    if max_amount is not None:
        query = query.filter(Transaction.amount <= max_amount)
    if search:
        like = f"%{search}%"
        query = query.filter(
            or_(
                Create_Account_Table.name.ilike(like),
                Create_Account_Table.email.ilike(like),
                Transaction.note.ilike(like),
            )
        )

    total = query.count()
    rows = (
        query.order_by(Transaction.transaction_datetime.desc())
        .offset(skip)
        .limit(limit)
        .all()
    )
    return schemas.AdminTransactionListResponse(
        total=total,
        transactions=[_to_admin_transaction(t) for t in rows],
    )


def update_transaction(
    db: Session,
    transaction_id: int,
    data: schemas.AdminTransactionUpdate,
) -> schemas.AdminTransactionOut:
    transaction = _get_transaction(db, transaction_id)
    update_data = data.model_dump(exclude_unset=True)
    for field, value in update_data.items():
        setattr(transaction, field, value)
    db.commit()
    db.refresh(transaction)
    return _to_admin_transaction(transaction)


def delete_transaction(db: Session, transaction_id: int) -> schemas.AdminTransactionOut:
    transaction = _get_transaction(db, transaction_id)
    result = _to_admin_transaction(transaction)
    db.delete(transaction)
    db.commit()
    return result


def suspicious_transactions(
    db: Session,
    threshold: Optional[float] = None,
    limit: int = 50,
) -> schemas.SuspiciousListResponse:
    amounts = [
        amount
        for (amount,) in db.query(Transaction.amount)
        .filter(Transaction.type == TransactionType.expense)
        .all()
        if amount is not None
    ]
    if not amounts:
        return schemas.SuspiciousListResponse(threshold=0, count=0, transactions=[])

    if threshold is None:
        threshold = (
            mean(amounts) + 2 * stdev(amounts)
            if len(amounts) > 1
            else float(amounts[0])
        )

    rows = (
        db.query(Transaction)
        .filter(
            Transaction.type == TransactionType.expense,
            Transaction.amount > threshold,
        )
        .order_by(Transaction.amount.desc())
        .limit(limit)
        .all()
    )
    return schemas.SuspiciousListResponse(
        threshold=round(threshold, 2),
        count=len(rows),
        transactions=[
            schemas.SuspiciousTransactionOut(
                transaction=_to_admin_transaction(t),
                reason=(
                    f"Expense amount exceeds anomaly threshold "
                    f"(mean + 2x std dev of all expenses)"
                ),
                threshold=round(threshold, 2),
            )
            for t in rows
        ],
    )


def category_analytics(db: Session) -> List[schemas.CategoryAnalyticsItem]:
    rows = (
        db.query(
            Transaction.category,
            func.count(Transaction.id).label("count"),
            func.sum(Transaction.amount).label("total_amount"),
            _income_expr(),
            _expense_expr(),
        )
        .group_by(Transaction.category)
        .order_by(func.sum(Transaction.amount).desc())
        .all()
    )
    return [
        schemas.CategoryAnalyticsItem(
            category=row.category,
            count=int(row.count or 0),
            total_amount=float(row.total_amount or 0),
            income_amount=float(row.total_income or 0),
            expense_amount=float(row.total_expense or 0),
        )
        for row in rows
    ]


# ------------------------------ Platform stats ------------------------------


def platform_stats(db: Session) -> schemas.PlatformStats:
    total_users = (
        db.query(func.count(Create_Account_Table.profile_id)).scalar() or 0
    )
    active_users = (
        db.query(func.count(Create_Account_Table.profile_id))
        .filter(Create_Account_Table.is_active.is_(True))
        .scalar()
        or 0
    )
    total_transactions = db.query(func.count(Transaction.id)).scalar() or 0
    total_income = float(
        db.query(_income_expr().label("total_income")).scalar() or 0
    )
    total_expense = float(
        db.query(_expense_expr().label("total_expense")).scalar() or 0
    )

    month_col = func.to_char(Create_Account_Table.created_at, "YYYY-MM").label("month")
    new_users = (
        db.query(month_col, func.count().label("num"))
        .group_by(month_col)
        .order_by(month_col)
        .all()
    )
    new_users_per_month = [
        schemas.MonthlyCount(month=row.month, count=row.num) for row in new_users
    ]

    tx_month = func.to_char(Transaction.transaction_datetime, "YYYY-MM").label("month")
    trends = (
        db.query(tx_month, func.count().label("num"))
        .group_by(tx_month)
        .order_by(tx_month)
        .all()
    )
    usage_trends = [
        schemas.MonthlyCount(month=row.month, count=row.num) for row in trends
    ]

    top_spender_rows = (
        db.query(
            Create_Account_Table.profile_id,
            Create_Account_Table.name,
            Create_Account_Table.email,
            _income_expr(),
            _expense_expr(),
            func.count(Transaction.id).label("transaction_count"),
        )
        .join(Transaction, Transaction.user_id == Create_Account_Table.profile_id)
        .group_by(
            Create_Account_Table.profile_id,
            Create_Account_Table.name,
            Create_Account_Table.email,
        )
        .order_by(_expense_expr().desc())
        .limit(10)
        .all()
    )
    top_spenders = [
        schemas.TopSpender(
            profile_id=row.profile_id,
            name=row.name,
            email=row.email,
            total_income=float(row.total_income or 0),
            total_expense=float(row.total_expense or 0),
            transaction_count=int(row.transaction_count or 0),
        )
        for row in top_spender_rows
    ]

    return schemas.PlatformStats(
        total_users=total_users,
        active_users=active_users,
        inactive_users=total_users - active_users,
        total_transactions=total_transactions,
        total_income=total_income,
        total_expense=total_expense,
        platform_balance=total_income - total_expense,
        avg_income_per_user=(total_income / total_users) if total_users else 0,
        avg_expense_per_user=(total_expense / total_users) if total_users else 0,
        new_users_per_month=new_users_per_month,
        usage_trends=usage_trends,
        top_spenders=top_spenders,
        top_categories=category_analytics(db)[:10],
    )


# ------------------------------ Category config ------------------------------


def list_categories(db: Session) -> List[Category]:
    return db.query(Category).order_by(Category.name).all()


def create_category(db: Session, name: str) -> Category:
    clean = name.strip()
    if not clean:
        raise HTTPException(status_code=400, detail="Category name cannot be empty")
    existing = (
        db.query(Category)
        .filter(func.lower(Category.name) == clean.lower())
        .first()
    )
    if existing:
        raise HTTPException(status_code=400, detail="Category already exists")
    category = Category(name=clean)
    db.add(category)
    db.commit()
    db.refresh(category)
    return category


def update_category(db: Session, category_id: int, data: schemas.CategoryUpdate) -> Category:
    category = db.query(Category).filter(Category.id == category_id).first()
    if not category:
        raise HTTPException(status_code=404, detail="Category not found")

    if data.name is not None:
        clean = data.name.strip()
        if not clean:
            raise HTTPException(status_code=400, detail="Category name cannot be empty")
        duplicate = (
            db.query(Category)
            .filter(
                func.lower(Category.name) == clean.lower(),
                Category.id != category_id,
            )
            .first()
        )
        if duplicate:
            raise HTTPException(status_code=400, detail="Category already exists")
        category.name = clean
    if data.is_active is not None:
        category.is_active = data.is_active

    db.commit()
    db.refresh(category)
    return category


def delete_category(db: Session, category_id: int) -> Category:
    category = db.query(Category).filter(Category.id == category_id).first()
    if not category:
        raise HTTPException(status_code=404, detail="Category not found")
    db.delete(category)
    db.commit()
    return category


# ---------------------------------- Export -----------------------------------


def _users_csv_rows(db: Session):
    users = _build_user_stats(db, db.query(Create_Account_Table).all())
    for u in users:
        yield [
            u.profile_id,
            u.name,
            u.email,
            u.mobile_number,
            u.role,
            u.is_active,
            u.created_at,
            u.total_income,
            u.total_expense,
            u.transaction_count,
        ]


def export_users_csv(db: Session) -> str:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(
        [
            "profile_id",
            "name",
            "email",
            "mobile_number",
            "role",
            "is_active",
            "created_at",
            "total_income",
            "total_expense",
            "transaction_count",
        ]
    )
    for row in _users_csv_rows(db):
        writer.writerow(row)
    return output.getvalue()


def export_users_xlsx(db: Session) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Users"
    ws.append(
        [
            "profile_id",
            "name",
            "email",
            "mobile_number",
            "role",
            "is_active",
            "created_at",
            "total_income",
            "total_expense",
            "transaction_count",
        ]
    )
    for row in _users_csv_rows(db):
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_transactions_csv(db: Session) -> str:
    rows = (
        db.query(Transaction)
        .join(Create_Account_Table, Transaction.user_id == Create_Account_Table.profile_id)
        .order_by(Transaction.transaction_datetime.desc())
        .all()
    )
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(
        [
            "id",
            "user_id",
            "user_name",
            "user_email",
            "type",
            "amount",
            "category",
            "note",
            "transaction_datetime",
            "created_at",
            "updated_at",
        ]
    )
    for t in rows:
        writer.writerow(
            [
                t.id,
                t.user_id,
                t.user.name,
                t.user.email,
                t.type.value,
                t.amount,
                t.category,
                t.note,
                t.transaction_datetime,
                t.created_at,
                t.updated_at,
            ]
        )
    return output.getvalue()


def export_transactions_xlsx(db: Session) -> bytes:
    rows = (
        db.query(Transaction)
        .join(Create_Account_Table, Transaction.user_id == Create_Account_Table.profile_id)
        .order_by(Transaction.transaction_datetime.desc())
        .all()
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(
        [
            "id",
            "user_id",
            "user_name",
            "user_email",
            "type",
            "amount",
            "category",
            "note",
            "transaction_datetime",
            "created_at",
            "updated_at",
        ]
    )
    for t in rows:
        ws.append(
            [
                t.id,
                t.user_id,
                t.user.name,
                t.user.email,
                t.type.value,
                t.amount,
                t.category,
                t.note,
                t.transaction_datetime,
                t.created_at,
                t.updated_at,
            ]
        )
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
